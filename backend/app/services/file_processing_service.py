import io
import os
import re
import time
import zipfile
import logging
from dataclasses import dataclass, field
from typing import List, Optional, Literal, Dict, Any
from html.parser import HTMLParser
from email import policy
from email.parser import BytesParser
import pymupdf as fitz
import docx

logger = logging.getLogger("shiftly.files")

MAX_FILE_SIZE_BYTES = 25 * 1024 * 1024  # 25 MB safety limit
MAX_ZIP_ENTRIES = 100  # Maximum archive entries allowed (Zip Bomb protection)
MAX_ZIP_UNCOMPRESSED_BYTES = 50 * 1024 * 1024  # 50 MB uncompressed limit
SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx", ".chat", ".log", ".csv", ".zip", ".eml", ".mbox"}


class FileProcessingError(Exception):
    """Base exception for file processing."""
    pass


class UnsupportedFileTypeError(FileProcessingError):
    """Raised when an uploaded file extension or content is unsupported."""
    pass


class FileOversizedError(FileProcessingError):
    """Raised when an uploaded file exceeds the configured size limit."""
    pass


class EmptyFileError(FileProcessingError):
    """Raised when an uploaded file is empty or contains no extractable text."""
    pass


class NoSelectableTextPDFError(FileProcessingError):
    """Raised when a PDF contains no selectable text (e.g., scanned/image-only)."""
    pass


class InvalidWhatsAppZipError(FileProcessingError):
    """Raised when a ZIP archive does not contain a valid WhatsApp chat transcript (_chat.txt)."""
    pass


class ZipSecurityError(FileProcessingError):
    """Raised when an uploaded ZIP file violates security constraints (Zip Slip / Path Traversal)."""
    pass


class ZipBombError(ZipSecurityError):
    """Raised when a ZIP exceeds entry count or uncompressed size thresholds."""
    pass


class EmailParsingError(FileProcessingError):
    """Raised when an email file (.eml) or archive (.mbox) is corrupted or cannot be parsed."""
    pass


@dataclass
class ContentBlock:
    text: str
    source_name: str
    source_type: Literal["Chat Export", "Email Thread", "Meeting Transcript", "Document"]
    location: str
    sender: Optional[str] = None


@dataclass
class ExtractedFileContent:
    filename: str
    source_type: Literal["Chat Export", "Email Thread", "Meeting Transcript", "Document"]
    full_text: str
    blocks: List[ContentBlock] = field(default_factory=list)


def parse_sender_from_line(line: str) -> Optional[str]:
    """Extracts speaker/sender name if matching common chat/transcript patterns."""
    # Pattern: [timestamp] Speaker: message OR Speaker: message OR timestamp - Speaker: message
    patterns = [
        r"^\[.*?\]\s*([^:]+?)\s*:",                 # [10/12/2024, 08:34] David Miller: ...
        r"^\d{1,2}:\d{2}(?::\d{2})?\s*-\s*([^:]+?)\s*:", # 08:34 - David Miller: ...
        r"^([A-Z][A-Za-z0-9\s.()-]{1,35})\s*:",     # David Miller (Client): ...
    ]
    for pat in patterns:
        m = re.match(pat, line.strip())
        if m:
            candidate = m.group(1).strip()
            # Ignore common false positives like "Note:", "Date:", "Time:"
            if candidate.lower() not in {"note", "date", "time", "http", "https", "subject"}:
                return candidate
    return None


def decode_text_file_bytes(file_bytes: bytes, filename: str) -> str:
    """
    Decodes text file bytes using BOM detection, Windows Notepad encodings,
    and safe fallbacks (UTF-8, UTF-8 with BOM, UTF-16 LE, UTF-16 BE, CP1252).
    Prevents corrupt latin-1 decoding of UTF-16 files.
    """
    if not file_bytes or not file_bytes.strip():
        raise EmptyFileError(f"The file '{filename}' is empty or contains no text.")

    # 1. Explicit BOM detection
    if file_bytes.startswith(b"\xef\xbb\xbf"):
        return file_bytes[3:].decode("utf-8", errors="replace")
    elif file_bytes.startswith(b"\xff\xfe\x00\x00"):
        return file_bytes[4:].decode("utf-32-le", errors="replace")
    elif file_bytes.startswith(b"\x00\x00\xfe\xff"):
        return file_bytes[4:].decode("utf-32-be", errors="replace")
    elif file_bytes.startswith(b"\xff\xfe"):
        return file_bytes[2:].decode("utf-16-le", errors="replace")
    elif file_bytes.startswith(b"\xfe\xff"):
        return file_bytes[2:].decode("utf-16-be", errors="replace")

    # 2. Heuristic for UTF-16 without BOM (common in Windows Notepad / PowerShell exports)
    # Detect high frequency of alternating null bytes
    if len(file_bytes) >= 2 and (file_bytes[1::2].count(b"\x00") > len(file_bytes) // 4):
        try:
            decoded = file_bytes.decode("utf-16-le")
            if "\x00" not in decoded:
                return decoded
        except UnicodeDecodeError:
            pass
    elif len(file_bytes) >= 2 and (file_bytes[0::2].count(b"\x00") > len(file_bytes) // 4):
        try:
            decoded = file_bytes.decode("utf-16-be")
            if "\x00" not in decoded:
                return decoded
        except UnicodeDecodeError:
            pass

    # 3. Standard UTF-8 without BOM (must not contain null bytes)
    try:
        decoded = file_bytes.decode("utf-8")
        if "\x00" not in decoded:
            return decoded
    except UnicodeDecodeError:
        pass

    # 4. Standard Windows-1252 (ANSI Windows default)
    try:
        decoded = file_bytes.decode("cp1252")
        if "\x00" not in decoded:
            return decoded
    except UnicodeDecodeError:
        pass

    # 5. Latin-1 fallback with null-byte sanitization
    try:
        decoded = file_bytes.decode("latin-1", errors="replace")
        cleaned = decoded.replace("\x00", "")
        if cleaned.strip():
            return cleaned
    except Exception:
        pass

    raise FileProcessingError(f"Could not decode text file '{filename}'. Please ensure it is saved in a supported text encoding (UTF-8, UTF-16, or ANSI).")


def extract_from_txt(file_bytes: bytes, filename: str) -> ExtractedFileContent:
    """Extracts text from TXT or chat export files, supporting multiple encodings."""
    text = decode_text_file_bytes(file_bytes, filename)
    if not text or not text.strip():
        raise EmptyFileError(f"The file '{filename}' is empty or contains no text.")

    # Determine if chat export or general document
    is_chat = any(parse_sender_from_line(line) for line in text.splitlines()[:20])
    source_type: Literal["Chat Export", "Document"] = "Chat Export" if is_chat else "Document"

    lines = [l for l in text.splitlines()]
    blocks: List[ContentBlock] = []
    full_text_lines: List[str] = []

    message_idx = 1
    for line_idx, raw_line in enumerate(lines, start=1):
        line = raw_line.strip()
        if not line:
            continue

        sender = parse_sender_from_line(line)
        location = f"Message #{message_idx}" if (is_chat and sender) else f"Line {line_idx}"
        if sender and is_chat:
            message_idx += 1

        blocks.append(
            ContentBlock(
                text=line,
                source_name=filename,
                source_type=source_type,
                location=location,
                sender=sender,
            )
        )
        full_text_lines.append(line)

    full_text = "\n".join(full_text_lines).strip()
    if not full_text:
        raise EmptyFileError(f"The file '{filename}' contains only whitespace or blank lines.")

    return ExtractedFileContent(
        filename=filename,
        source_type=source_type,
        full_text=full_text,
        blocks=blocks,
    )


def extract_from_pdf(file_bytes: bytes, filename: str) -> ExtractedFileContent:
    """Extracts text page-by-page from PDF using PyMuPDF."""
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as e:
        raise FileProcessingError(f"Failed to open PDF '{filename}': {str(e)}") from e

    try:
        if len(doc) == 0:
            raise EmptyFileError(f"The PDF '{filename}' contains 0 pages.")

        blocks: List[ContentBlock] = []
        full_text_parts: List[str] = []
        total_extracted_chars = 0

        for page_idx in range(len(doc)):
            page = doc[page_idx]
            page_text = page.get_text("text").strip()
            page_num = page_idx + 1

            if page_text:
                total_extracted_chars += len(page_text)
                for line in page_text.splitlines():
                    clean_line = line.strip()
                    if clean_line:
                        sender = parse_sender_from_line(clean_line)
                        blocks.append(
                            ContentBlock(
                                text=clean_line,
                                source_name=filename,
                                source_type="Document",
                                location=f"Page {page_num}",
                                sender=sender,
                            )
                        )
                full_text_parts.append(f"[Page {page_num}]\n{page_text}")

        if total_extracted_chars == 0:
            raise NoSelectableTextPDFError(
                "This PDF does not contain selectable text. Scanned PDF/OCR support is not available yet."
            )

        full_text = "\n\n".join(full_text_parts).strip()
        return ExtractedFileContent(
            filename=filename,
            source_type="Document",
            full_text=full_text,
            blocks=blocks,
        )
    finally:
        doc.close()


def extract_from_docx(file_bytes: bytes, filename: str) -> ExtractedFileContent:
    """Extracts text paragraph-by-paragraph from DOCX using python-docx."""
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise FileProcessingError(f"Failed to read Word document '{filename}': {str(e)}") from e

    blocks: List[ContentBlock] = []
    full_text_parts: List[str] = []
    current_section: Optional[str] = None
    para_count = 0

    for p in doc.paragraphs:
        p_text = p.text.strip()
        if not p_text:
            continue

        para_count += 1
        # Check if heading
        if p.style and p.style.name and p.style.name.startswith("Heading"):
            current_section = p_text

        location = f"Section '{current_section}' (Para {para_count})" if current_section else f"Paragraph {para_count}"
        sender = parse_sender_from_line(p_text)

        blocks.append(
            ContentBlock(
                text=p_text,
                source_name=filename,
                source_type="Document",
                location=location,
                sender=sender,
            )
        )
        full_text_parts.append(p_text)

    # Also extract tables if present
    for table_idx, table in enumerate(doc.tables, start=1):
        for row_idx, row in enumerate(table.rows, start=1):
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                blocks.append(
                    ContentBlock(
                        text=row_text,
                        source_name=filename,
                        source_type="Document",
                        location=f"Table {table_idx}, Row {row_idx}",
                    )
                )
                full_text_parts.append(row_text)

    if not full_text_parts:
        raise EmptyFileError(f"The document '{filename}' contains no readable text.")

    full_text = "\n\n".join(full_text_parts).strip()
    return ExtractedFileContent(
        filename=filename,
        source_type="Document",
        full_text=full_text,
        blocks=blocks,
    )


# Safe HTML Text Extractor for HTML emails
class SafeHTMLTextExtractor(HTMLParser):
    """Safely extracts visible text from HTML emails without executing scripts or loading assets."""
    def __init__(self):
        super().__init__()
        self._chunks: List[str] = []
        self._ignore_stack: int = 0

    def handle_starttag(self, tag: str, attrs):
        tag_lower = tag.lower()
        if tag_lower in ("script", "style", "head", "meta", "noscript"):
            self._ignore_stack += 1
        elif tag_lower in ("p", "div", "br", "h1", "h2", "h3", "h4", "h5", "h6", "tr", "li"):
            self._chunks.append("\n")

    def handle_endtag(self, tag: str):
        tag_lower = tag.lower()
        if tag_lower in ("script", "style", "head", "meta", "noscript"):
            if self._ignore_stack > 0:
                self._ignore_stack -= 1
        elif tag_lower in ("p", "div", "h1", "h2", "h3", "h4", "h5", "h6", "tr", "li"):
            self._chunks.append("\n")

    def handle_data(self, data: str):
        if self._ignore_stack == 0:
            self._chunks.append(data)

    def get_text(self) -> str:
        raw = "".join(self._chunks)
        cleaned = re.sub(r"\n{3,}", "\n\n", raw)
        return cleaned.strip()


def extract_text_from_html(html_str: str) -> str:
    extractor = SafeHTMLTextExtractor()
    extractor.feed(html_str)
    return extractor.get_text()


# WhatsApp Parsing Patterns
WHATSAPP_IOS_PATTERN = re.compile(
    r"^\[(\d{1,2}[./-]\d{1,2}[./-]\d{2,4}),?\s+(\d{1,2}:\d{2}(?::\d{2})?(?:\s*[APap][Mm])?)\]\s*(?:([^:]+?):\s*(.*))?$"
)
WHATSAPP_ANDROID_PATTERN = re.compile(
    r"^(\d{1,2}[./-]\d{1,2}[./-]\d{2,4}),?\s+(\d{1,2}:\d{2}(?::\d{2})?(?:\s*[APap][Mm])?)\s*-\s*(?:([^:]+?):\s*(.*))?$"
)

MEDIA_OMITTED_PATTERNS = [
    re.compile(r"<Media omitted>", re.IGNORECASE),
    re.compile(r"image omitted", re.IGNORECASE),
    re.compile(r"video omitted", re.IGNORECASE),
    re.compile(r"audio omitted", re.IGNORECASE),
    re.compile(r"sticker omitted", re.IGNORECASE),
    re.compile(r"document omitted", re.IGNORECASE),
    re.compile(r"Contact card omitted", re.IGNORECASE),
    re.compile(r"\(file attached\)", re.IGNORECASE),
]


def sanitize_whatsapp_text(raw_text: str) -> str:
    """Strips invisible Unicode control characters and normalizes spaces."""
    cleaned = (
        raw_text.replace("\u200e", "")
        .replace("\u200f", "")
        .replace("\u202a", "")
        .replace("\u202c", "")
        .replace("\u202e", "")
    )
    cleaned = cleaned.replace("\u202f", " ").replace("\xa0", " ")
    return cleaned


def parse_whatsapp_chat(text: str, filename: str) -> ExtractedFileContent:
    """
    Parses WhatsApp chat export text into structured messages, supporting
    both iOS bracketed and Android dash timestamp formats, multiline messages,
    Unicode/Hinglish/emojis, and filtering media placeholders.
    """
    sanitized = sanitize_whatsapp_text(text)
    raw_lines = sanitized.splitlines()

    parsed_messages: List[Dict[str, Any]] = []
    current_msg: Optional[Dict[str, Any]] = None

    for raw_line in raw_lines:
        line = raw_line.strip()
        if not line:
            if current_msg:
                current_msg["body"] += "\n"
            continue

        # Match iOS bracketed or Android dash format
        m_ios = WHATSAPP_IOS_PATTERN.match(line)
        m_android = None if m_ios else WHATSAPP_ANDROID_PATTERN.match(line)
        m = m_ios or m_android

        if m:
            date_part = m.group(1).strip()
            time_part = m.group(2).strip()
            sender_candidate = m.group(3)
            body_candidate = m.group(4) if m.group(4) is not None else ""

            if sender_candidate:
                sender = sender_candidate.strip()
                clean_body = body_candidate.strip()
                for pat in MEDIA_OMITTED_PATTERNS:
                    clean_body = pat.sub("[Media attachment omitted]", clean_body)

                if current_msg:
                    parsed_messages.append(current_msg)

                current_msg = {
                    "date": date_part,
                    "time": time_part,
                    "sender": sender,
                    "body": clean_body,
                }
                continue
            else:
                # System notification (e.g. encryption notice or group change)
                if "end-to-end encrypted" in line.lower():
                    continue
                if current_msg:
                    parsed_messages.append(current_msg)
                    current_msg = None
                parsed_messages.append({
                    "date": date_part,
                    "time": time_part,
                    "sender": "System",
                    "body": line,
                })
                continue

        # Multiline message continuation
        if current_msg:
            clean_line = line
            for pat in MEDIA_OMITTED_PATTERNS:
                clean_line = pat.sub("[Media attachment omitted]", clean_line)
            current_msg["body"] += "\n" + clean_line
        else:
            # Leading non-timestamped line
            if "end-to-end encrypted" in line.lower():
                continue
            current_msg = {
                "date": "",
                "time": "",
                "sender": "System",
                "body": line,
            }

    if current_msg:
        parsed_messages.append(current_msg)

    if not parsed_messages:
        raise EmptyFileError(f"WhatsApp chat in '{filename}' contains no readable messages.")

    blocks: List[ContentBlock] = []
    full_text_lines: List[str] = []

    for idx, msg in enumerate(parsed_messages, start=1):
        timestamp = f"{msg['date']} {msg['time']}".strip()
        sender = msg["sender"]
        body = msg["body"].strip()
        if not body:
            continue

        location = f"Message #{idx} ({timestamp})" if timestamp else f"Message #{idx}"
        blocks.append(
            ContentBlock(
                text=f"{sender}: {body}",
                source_name=filename,
                source_type="Chat Export",
                location=location,
                sender=sender if sender != "System" else None,
            )
        )
        if timestamp:
            full_text_lines.append(f"[{timestamp}] {sender}: {body}")
        else:
            full_text_lines.append(f"{sender}: {body}")

    full_text = "\n\n".join(full_text_lines).strip()
    if not full_text:
        raise EmptyFileError(f"WhatsApp chat in '{filename}' contains no readable message text.")

    return ExtractedFileContent(
        filename=filename,
        source_type="Chat Export",
        full_text=full_text,
        blocks=blocks,
    )


def extract_from_whatsapp_zip(file_bytes: bytes, filename: str) -> ExtractedFileContent:
    """
    Safely inspects a WhatsApp export ZIP archive in memory.
    Enforces zip bomb bounds (max 100 entries, max 50 MB uncompressed),
    strictly validates against Zip Slip / path traversal, ignores all media binaries,
    and extracts only the chat transcript (_chat.txt).
    """
    try:
        zf = zipfile.ZipFile(io.BytesIO(file_bytes))
    except zipfile.BadZipFile as bzf:
        raise UnsupportedFileTypeError(f"File '{filename}' is not a valid or readable ZIP archive: {str(bzf)}")
    except Exception as e:
        raise FileProcessingError(f"Failed to inspect ZIP archive '{filename}': {str(e)}")

    with zf:
        infolist = zf.infolist()
        # 1. Entry count limit (Zip Bomb protection)
        if len(infolist) > MAX_ZIP_ENTRIES:
            raise ZipBombError(
                f"ZIP archive contains {len(infolist)} entries, exceeding the maximum limit of {MAX_ZIP_ENTRIES} entries."
            )

        # 2. Total uncompressed size limit (Zip Bomb protection)
        total_uncompressed = sum(info.file_size for info in infolist)
        if total_uncompressed > MAX_ZIP_UNCOMPRESSED_BYTES:
            max_mb = MAX_ZIP_UNCOMPRESSED_BYTES // (1024 * 1024)
            raise ZipBombError(
                f"ZIP archive uncompressed size ({total_uncompressed // (1024 * 1024)} MB) exceeds allowed limit of {max_mb} MB."
            )

        # 3. Path traversal / Zip Slip validation
        for info in infolist:
            entry_name = info.filename.replace("\\", "/")
            if os.path.isabs(entry_name) or entry_name.startswith("/") or bool(re.match(r"^[a-zA-Z]:", entry_name)):
                raise ZipSecurityError(f"ZIP archive contains unsafe absolute path: '{info.filename}'")
            segments = entry_name.split("/")
            if ".." in segments:
                raise ZipSecurityError(f"ZIP archive contains directory traversal sequence (..): '{info.filename}'")

        # 4. Locate WhatsApp chat text file
        target_member: Optional[zipfile.ZipInfo] = None
        txt_members: List[zipfile.ZipInfo] = []

        for info in infolist:
            if info.is_dir():
                continue
            base = os.path.basename(info.filename)
            if base.lower() == "_chat.txt":
                target_member = info
                break
            elif re.search(r"whatsapp.*chat.*\.txt$", base, re.IGNORECASE):
                target_member = info
                break
            elif base.lower().endswith(".txt"):
                txt_members.append(info)

        if not target_member and len(txt_members) == 1:
            target_member = txt_members[0]

        if not target_member:
            raise InvalidWhatsAppZipError(
                "ZIP does not contain a readable WhatsApp chat transcript (_chat.txt). "
                "Please export your WhatsApp chat (with or without media) and upload the resulting .zip file."
            )

        # 5. Read ONLY the chat transcript in-memory (ignore all media binaries)
        try:
            chat_bytes = zf.read(target_member)
        except Exception as e:
            raise FileProcessingError(f"Failed to extract chat transcript from ZIP: {str(e)}")

    # 6. Decode and parse
    text = decode_text_file_bytes(chat_bytes, target_member.filename)
    return parse_whatsapp_chat(text, filename)


def extract_from_eml(file_bytes: bytes, filename: str) -> ExtractedFileContent:
    """Extracts headers, body text, and source blocks from an RFC 822/2822 email file (.eml)."""
    try:
        msg = BytesParser(policy=policy.default).parsebytes(file_bytes)
    except Exception as e:
        raise EmailParsingError(f"Failed to parse email file '{filename}': {str(e)}")

    subject = str(msg.get("Subject", "")).strip() or "(No Subject)"
    from_hdr = str(msg.get("From", "")).strip() or "Unknown Sender"
    to_hdr = str(msg.get("To", "")).strip()
    cc_hdr = str(msg.get("Cc", "")).strip()
    date_hdr = str(msg.get("Date", "")).strip()

    body_text = ""
    try:
        body_part = msg.get_body(preferencelist=("plain",))
        if body_part:
            body_text = body_part.get_content()
        else:
            html_part = msg.get_body(preferencelist=("html",))
            if html_part:
                html_content = html_part.get_content()
                body_text = extract_text_from_html(html_content)
    except Exception:
        pass

    if not body_text or not body_text.strip():
        text_chunks = []
        for part in msg.walk():
            content_type = part.get_content_type()
            if content_type == "text/plain":
                try:
                    payload = part.get_content()
                    if payload and payload.strip():
                        text_chunks.append(payload.strip())
                except Exception:
                    pass
            elif content_type == "text/html" and not text_chunks:
                try:
                    payload = part.get_content()
                    if payload and payload.strip():
                        text_chunks.append(extract_text_from_html(payload))
                except Exception:
                    pass
        if text_chunks:
            body_text = "\n\n".join(text_chunks)

    if not body_text or not body_text.strip():
        raise EmptyFileError(f"The email '{filename}' contains no readable text content.")

    body_text = body_text.strip()

    header_lines = []
    if from_hdr:
        header_lines.append(f"From: {from_hdr}")
    if to_hdr:
        header_lines.append(f"To: {to_hdr}")
    if cc_hdr:
        header_lines.append(f"Cc: {cc_hdr}")
    if date_hdr:
        header_lines.append(f"Date: {date_hdr}")
    if subject:
        header_lines.append(f"Subject: {subject}")

    header_str = "\n".join(header_lines)
    full_text = f"{header_str}\n\n{body_text}"

    blocks: List[ContentBlock] = []
    blocks.append(
        ContentBlock(
            text=header_str,
            source_name=filename,
            source_type="Email Thread",
            location="Email Header",
            sender=from_hdr,
        )
    )

    paras = [p.strip() for p in body_text.split("\n\n") if p.strip()]
    for idx, p in enumerate(paras, start=1):
        blocks.append(
            ContentBlock(
                text=p,
                source_name=filename,
                source_type="Email Thread",
                location=f"Email Body (Section {idx})",
                sender=from_hdr,
            )
        )

    return ExtractedFileContent(
        filename=filename,
        source_type="Email Thread",
        full_text=full_text,
        blocks=blocks,
    )


def extract_from_mbox(file_bytes: bytes, filename: str) -> ExtractedFileContent:
    """Extracts messages from an email archive (.mbox) safely in-memory."""
    if not file_bytes or not file_bytes.strip():
        raise EmptyFileError(f"The mbox file '{filename}' is empty.")

    lines = file_bytes.splitlines(keepends=True)
    msg_chunks: List[bytes] = []
    current_chunk: List[bytes] = []

    for line in lines:
        if line.startswith(b"From ") and current_chunk:
            msg_chunks.append(b"".join(current_chunk))
            current_chunk = []
            if len(msg_chunks) >= 50:
                break
        current_chunk.append(line)

    if current_chunk and len(msg_chunks) < 50:
        msg_chunks.append(b"".join(current_chunk))

    if not msg_chunks:
        raise EmptyFileError(f"The mbox file '{filename}' contains no messages.")

    blocks: List[ContentBlock] = []
    full_text_messages: List[str] = []

    for idx, chunk in enumerate(msg_chunks, start=1):
        try:
            msg = BytesParser(policy=policy.default).parsebytes(chunk)
            subject = str(msg.get("Subject", "")).strip() or "(No Subject)"
            from_hdr = str(msg.get("From", "")).strip() or "Unknown Sender"
            date_hdr = str(msg.get("Date", "")).strip()
            to_hdr = str(msg.get("To", "")).strip()

            body_text = ""
            body_part = msg.get_body(preferencelist=("plain",))
            if body_part:
                body_text = body_part.get_content()
            else:
                html_part = msg.get_body(preferencelist=("html",))
                if html_part:
                    body_text = extract_text_from_html(html_part.get_content())

            if not body_text:
                for part in msg.walk():
                    if part.get_content_type() == "text/plain":
                        payload = part.get_content()
                        if payload:
                            body_text = payload
                            break

            clean_body = body_text.strip() if body_text else "[Empty message body]"
            msg_repr = f"Message #{idx}\nFrom: {from_hdr}\nTo: {to_hdr}\nDate: {date_hdr}\nSubject: {subject}\n\n{clean_body}"
            full_text_messages.append(msg_repr)

            blocks.append(
                ContentBlock(
                    text=msg_repr,
                    source_name=filename,
                    source_type="Email Thread",
                    location=f"Message #{idx} ({date_hdr})" if date_hdr else f"Message #{idx}",
                    sender=from_hdr,
                )
            )
        except Exception as e:
            logger.warning("Failed to parse mbox message #%d: %s", idx, e)
            continue

    if not full_text_messages:
        raise EmailParsingError(f"Could not extract readable messages from mbox '{filename}'.")

    full_text = "\n\n---\n\n".join(full_text_messages)
    return ExtractedFileContent(
        filename=filename,
        source_type="Email Thread",
        full_text=full_text,
        blocks=blocks,
    )


def process_uploaded_file(file_bytes: bytes, filename: str) -> ExtractedFileContent:
    """
    Validates file size, extension, and content headers, then routes to the appropriate parser.
    Sanitizes filenames to prevent path traversal.
    Returns normalized text and granular location ContentBlocks.
    """
    if not filename or not filename.strip():
        raise UnsupportedFileTypeError("Missing filename in uploaded file.")

    # Sanitize filename: extract basename, strip null bytes and directory traversal characters
    clean_filename = os.path.basename(filename.replace("\\", "/")).strip()
    clean_filename = re.sub(r"[\x00-\x1f\x7f]", "", clean_filename)
    if not clean_filename or clean_filename in (".", ".."):
        raise UnsupportedFileTypeError("Invalid or unsafe filename.")

    start_t = time.perf_counter()
    logger.info("file_processing_started filename=%s size_bytes=%d", clean_filename, len(file_bytes))

    try:
        if len(file_bytes) == 0:
            raise EmptyFileError(f"Uploaded file '{clean_filename}' is empty (0 bytes).")

        if len(file_bytes) > MAX_FILE_SIZE_BYTES:
            max_mb = MAX_FILE_SIZE_BYTES // (1024 * 1024)
            raise FileOversizedError(f"File '{clean_filename}' exceeds maximum allowed size of {max_mb} MB.")

        _, ext = os.path.splitext(clean_filename.lower())
        if ext not in SUPPORTED_EXTENSIONS:
            supported_str = ", ".join(sorted(SUPPORTED_EXTENSIONS))
            raise UnsupportedFileTypeError(
                f"Unsupported file format '{ext}'. Supported formats: {supported_str}"
            )

        if ext in {".txt", ".chat", ".log", ".csv"}:
            result = extract_from_txt(file_bytes, clean_filename)
        elif ext == ".pdf":
            if not file_bytes.startswith(b"%PDF-"):
                raise UnsupportedFileTypeError(f"File '{clean_filename}' is not a valid PDF document (missing PDF signature).")
            result = extract_from_pdf(file_bytes, clean_filename)
        elif ext == ".docx":
            # Standard DOCX files are zip containers starting with 'PK\x03\x04'
            if not file_bytes.startswith(b"PK\x03\x04"):
                raise UnsupportedFileTypeError(f"File '{clean_filename}' is not a valid Word document (missing DOCX/ZIP signature).")
            result = extract_from_docx(file_bytes, clean_filename)
        elif ext == ".zip":
            if not file_bytes.startswith(b"PK\x03\x04"):
                raise UnsupportedFileTypeError(f"File '{clean_filename}' is not a valid ZIP archive (missing ZIP signature).")
            result = extract_from_whatsapp_zip(file_bytes, clean_filename)
        elif ext == ".eml":
            result = extract_from_eml(file_bytes, clean_filename)
        elif ext == ".mbox":
            result = extract_from_mbox(file_bytes, clean_filename)
        else:
            raise UnsupportedFileTypeError(f"Unsupported file format '{ext}'")

        duration_ms = (time.perf_counter() - start_t) * 1000.0
        logger.info(
            "file_processing_completed filename=%s text_len=%d blocks=%d duration_ms=%.2f",
            clean_filename,
            len(result.full_text),
            len(result.blocks),
            duration_ms,
        )
        return result
    except Exception as exc:
        duration_ms = (time.perf_counter() - start_t) * 1000.0
        logger.error(
            "file_processing_failed filename=%s error=%s duration_ms=%.2f",
            clean_filename,
            type(exc).__name__,
            duration_ms,
        )
        raise

